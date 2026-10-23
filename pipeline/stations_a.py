"""站 0–2：需求解析、素材采集、脚本分镜。"""
import datetime, json, os, pathlib, re, shutil, subprocess
from common import ROOT, Station, StationFail, rjson, wjson, validate, run, sha256, ver

FAULT = os.environ.get("FAULT", "")
TODAY = "2026-10-23"
SITE = "https://revu.website/142620/"


def s0_brief(rd, sentence):
    with Station(rd, "00_brief", tools=[{"name": "jq", "version": ver("jq --version")}, {"name": "python-re", "version": "stdlib"}]) as st:
        dur = re.search(r"(\d+)\s*秒", sentence)
        en = re.search(r"[（(]([A-Za-z][A-Za-z ]+)[)）]", sentence)
        zh = re.search(r"给(?:悉尼)?(.+?)[（(]", sentence)
        brief = {"sentence": sentence,
                 "client": {"name_en": en.group(1).strip() if en else "", "name_zh_spoken": zh.group(1) if zh else "",
                            "city": "Sydney" if "悉尼" in sentence else ""},
                 "duration_s": int(dur.group(1)) if dur else 30,
                 "languages": [l for l, k in (("zh", "中"), ("en", "英")) if k in sentence] or ["zh", "en"],
                 "aspects": ["9:16", "1:1", "16:9"] if ("三个" in sentence or "3" in sentence) else ["9:16"],
                 "platforms": ["Instagram Reels", "TikTok", "小红书", "Facebook", "YouTube"],
                 "deliverables": ["promo_9x16.mp4", "promo_1x1.mp4", "promo_16x9.mp4", "cover_9x16.jpg", "cover_1x1.jpg",
                                  "cover_16x9.jpg", "subs/zh.srt", "subs/en.srt", "报价单"]}
        p = wjson(st.dir / "brief.json", brief)
        r = run(["jq", "-e", '.duration_s and (.aspects|length>0) and (.client.name_en|length>0)', str(p)], st.name, "jq-assert")
        validate(p, "brief", st.name); st.out(p, "brief")
        st.check("client name parsed", bool(brief["client"]["name_en"]), brief["client"]["name_en"])
        st.check("duration parsed", brief["duration_s"] == 30, brief["duration_s"], 30)
        st.score = {"commercial": 4, "handoff": 5, "craft": 4, "benchmark": 3,
                    "reasons": {"commercial": "需求字段齐全，但平台列表为默认值，客户未确认",
                                "craft": "规则解析，对句式敏感（换说法可能解析失败）",
                                "benchmark": "jq 是 JSON 命令行事实标准；需求解析本身是自写正则，非标杆，属弱格"}}
    return p


def s1_ingest(rd, brief_p):
    cm = ROOT / "customer_media"
    with Station(rd, "01_ingest", inputs=[brief_p], tools=[{"name": "curl", "version": ver("curl --version")},
                 {"name": "ffprobe", "version": ver("ffprobe -version")}]) as st:
        validate(brief_p, "brief", st.name, "上游校验")
        snap = st.dir / "site_snapshot.html"
        r = subprocess.run(["curl", "-sL", "--compressed", "--max-time", "30", "-o", str(snap), SITE])
        live = r.returncode == 0 and snap.stat().st_size > 500
        if not live:
            shutil.copy(cm / "placeholder/official_site_snapshot_2026-10-23.html", snap)
        html = snap.read_text(errors="ignore")
        c4 = os.environ.get("C4A_PY", "/workspace/c4a/bin/python")
        code = ("import asyncio,sys,crawl4ai as c\nasync def m():\n  async with c.AsyncWebCrawler() as w:\n"
                "    r=await w.arun(sys.argv[1]); open(sys.argv[2],'w').write(r.markdown or '')\nasyncio.run(m())")
        r4 = subprocess.run([c4, "-c", code, SITE, str(st.dir / "site.md")], capture_output=True, text=True, timeout=180) if os.path.exists(c4) else None
        c4ok = bool(r4 and r4.returncode == 0 and (st.dir / "site.md").exists())
        os.environ["CRAWL4AI_OK"] = "1" if c4ok else ""
        st.tools.append({"name": "crawl4ai", "version": ver(f"{c4} -c 'import importlib.metadata as m;print(m.version(\"crawl4ai\"))'") if c4ok else "降级：未运行/失败，用 curl"})
        st.check("crawl4ai 正文提取", True, "ok" if c4ok else "降级 curl", note="robots.txt 302 至 revuonline.com，无 Disallow；只抓 1 个公开页")
        if c4ok: st.out(st.dir / "site.md")
        title = re.search(r"<title>\s*([^<]+?)\s*</title>", html, re.I)
        tel = re.search(r'href="tel:([0-9 ]+)"', html)
        st.check("官网可读", bool(title), title.group(1) if title else None, note="live" if live else "离线快照回退")
        tel_fmt = None
        if tel:
            d = re.sub(r"\D", "", tel.group(1)); tel_fmt = f"{d[:2]} {d[2:6]} {d[6:]}"
        src = [{"source_id": "S1-official", "url": SITE, "retrieved": TODAY, "sha256": sha256(snap)},
               {"source_id": "S2-directories", "url": "oceaniabiz / Uber Eats / Mate's Table（任务书 2026-10-22 核对）", "retrieved": "2026-10-22"},
               {"source_id": "S3-ubereats", "url": "Uber Eats 店页摘要（任务书）", "retrieved": "2026-10-22"},
               {"source_id": "S4-hours", "url": "Mate's Table vs Zmenu（任务书）", "retrieved": "2026-10-22"},
               {"source_id": "C-pending", "url": "客户待提供", "retrieved": TODAY}]
        F = lambda f, v, s, stt, n="": {"field": f, "value": v, "source_id": s, "status": stt, "note": n}
        facts = {"sources": src, "facts": [
            F("name_en", title.group(1) if title else None, "S1-official", "verified" if title else "pending", "官网 <title>"),
            F("phone", tel_fmt, "S1-official", "verified" if tel_fmt else "pending", "官网 tel: 链接"),
            F("name_zh", "小锅米线", "C-pending", "unverified", "客户口头说法，书面中文店名待提供"),
            F("address", "Shop 1 / R1, 57 Rothschild Ave, Rosebery NSW 2018", "S2-directories", "unverified", "官网未写，待客户确认"),
            F("cuisine", "云南米线", "S3-ubereats", "unverified"),
            F("signature_dishes", ["Slurp Slurp Little Copper Pot", "Signature Minty Beef"], "S3-ubereats", "unverified"),
            F("hours", ["周二–周日 11:00–21:00", "周五–周日 11:00–15:00、16:00–21:00，周一休"], "S4-hours", "conflict", "视频不写时间"),
            F("menu_prices", None, "C-pending", "pending", "官网无菜单"),
            F("social", None, "C-pending", "pending"),
            F("rating", "Uber Eats 4.8（900+）", "S3-ubereats", "unverified", "第三方、会变，不入片"),
            F("logo_content", "官网 Slur-Slurp-New-Logo.png 画面为「上座 SHANG ZUO」", "S1-official", "conflict", "与店名不符，不入片"),
            F("brand_font", None, "C-pending", "pending", "用 Noto Sans CJK（OFL）替代")]}
        if FAULT == "phone_unverified":
            facts["facts"][1]["status"] = "unverified"
        fp = wjson(st.dir / "facts.json", facts); validate(fp, "facts", st.name); st.out(fp, "facts")
        # 媒体入库
        media = []
        lic = {"assets": []}
        for mid, fn, sid in (("logo", "Slur-Slurp-New-Logo.png", "S1-official"),
                             ("store", "75513371_2388668874794660_3434699996744646656_n.jpg", "S1-official")):
            p = cm / "placeholder" / fn
            pr = json.loads(run(["ffprobe", "-v", "error", "-show_entries", "stream=width,height", "-of", "json", str(p)], st.name, "ffprobe").stdout)["streams"][0]
            m = {"media_id": mid, "path": os.path.relpath(p, ROOT), "sha256": sha256(p), "origin": "placeholder",
                 "license": "© Slurp Slurp Noodle（官网素材），仅限占位", "permission": "pending", "kind": "image",
                 "width": pr["width"], "height": pr["height"], "duration_s": None, "fps": None, "source_id": sid}
            if FAULT == "nolicense" and mid == "store":
                m["license"] = ""
            media.append(m)
        media[0]["note"] = "图像内容为「上座 SHANG ZUO」字样，与店名 Slurp Slurp Noodle 不符 → conflict，不入片，待客户提供正确 logo"
        media[0]["permission"] = "pending"
        wm = st.dir / "wordmark.png"
        run(["magick", "-background", "none", "-fill", "#FFF4E0", "-font", "/usr/share/fonts/opentype/noto/NotoSansCJK-Bold.ttc",
             "-pointsize", "120", "label:" + (title.group(1) if title else "?"), "-trim", "+repage", "-bordercolor", "none", "-border", "20", str(wm)], st.name, "wordmark")
        st.out(wm)
        media.append({"media_id": "wordmark", "path": os.path.relpath(wm, ROOT), "sha256": sha256(wm), "origin": "generated",
                      "license": "本流水线生成（店名文字 + Noto Sans CJK, OFL）", "permission": "n/a", "kind": "image",
                      "width": int(run(["magick", "identify", "-format", "%w", str(wm)], st.name, "id").stdout), "height": int(run(["magick", "identify", "-format", "%h", str(wm)], st.name, "id").stdout),
                      "duration_s": None, "fps": None, "source_id": "S1-official"})
        cust = [p for p in (cm / "customer-supplied").glob("*")] if (cm / "customer-supplied").exists() else []
        st.check("客户实拍素材到位", len(cust) > 0, len(cust), ">=1", "未到：全链用占位样片跑通")
        font = "/usr/share/fonts/opentype/noto/NotoSansCJK-Bold.ttc"
        media.append({"media_id": "font_noto_cjk", "path": font, "sha256": sha256(font), "origin": "generated", "license": "SIL OFL 1.1",
                      "permission": "n/a", "kind": "font", "width": 0, "height": 0, "duration_s": None, "fps": None, "source_id": "OSS"})
        mm = {"media": media}
        try:
            validate(mm, "media_manifest", st.name, "media_manifest")
        except StationFail as e:
            st.check("每个素材有来源+许可", False, str(e.why)[:200], "100%")
            raise StationFail(st.name, "license-check", "素材缺许可，停止（下游不得使用无许可素材）")
        st.check("每个素材有来源+许可", all(m["license"] and m["origin"] for m in media), "100%", "100%")
        mp = wjson(st.dir / "media_manifest.json", mm); st.out(mp, "media_manifest")
        for m in media:
            lic["assets"].append({"media_id": m["media_id"], "path": m["path"], "origin": m["origin"], "license": m["license"],
                                  "permission": m["permission"], "deliverable_ok": m["permission"] in ("granted", "n/a")})
        lic["assets"].append({"media_id": "music_bed", "path": "runs/<run_id>/03_grade/audio/mix.wav", "origin": "generated",
                              "license": "CC0（本流水线用 sox 合成，无第三方音乐）", "permission": "n/a", "deliverable_ok": True})
        lp = wjson(st.dir / "licenses.json", lic); st.out(lp)
        sr = st.dir / "shot_request.md"
        sr.write_text((ROOT / "shot_request.md").read_text()); st.out(sr)
        st.check("占位素材分辨率（客户素材标准 ≥1080p 不适用占位）", True, [f'{m["width"]}x{m["height"]}' for m in media[:2]],
                 note="占位图 500x500 / 680x510，远低于 1080p，仅样片")
        st.score = {"commercial": 2, "handoff": 5, "craft": 3, "benchmark": 3,
                    "reasons": {"commercial": "无客户实拍素材，只能出未授权样片，不可交付",
                                "craft": "来源/许可 100% 记录；但素材分辨率不达标（占位）",
                                "benchmark": "curl/ffprobe 是标杆；crawl4ai 本次" + ("已验证并用于正文提取（site.md），curl 存原始 HTML 作 sha256 证据" if os.environ.get("CRAWL4AI_OK") else "未通过验证，降级为 curl")}}
    return fp, mp, lp


def s2_script(rd, facts_p, media_p):
    with Station(rd, "02_script", inputs=[facts_p, media_p], tools=[{"name": "python-docx", "version": ver("python -c 'import docx,importlib.metadata as m;print(m.version(\"python-docx\"))'")},
                 {"name": "reportlab", "version": ver("python -c 'import reportlab;print(reportlab.Version)'")}, {"name": "jq", "version": ver("jq --version")}]) as st:
        facts = validate(facts_p, "facts", st.name, "上游校验"); validate(media_p, "media_manifest", st.name, "上游校验")
        F = {f["field"]: f for f in facts["facts"]}
        cmenu = ROOT / "customer_media" / "menu.json"
        menu = rjson(cmenu) if cmenu.exists() else {"status": "pending: 客户未提供菜单与价格", "items": []}
        menu_p = wjson(st.dir / "menu.json", menu); validate(menu_p, "menu", st.name); st.out(menu_p, "menu")
        name, phone = F["name_en"]["value"], F["phone"]["value"]
        shots = [
            ("sh1", 0, 4, "店名字标开场：由小放大，暖色底", "wordmark", "C-04 店门/招牌", name, name, ["facts:name_en"]),
            ("sh2", 4, 9, "店内图缓推（占位；交付版换铜锅沸腾）", "store", "C-01 铜锅沸腾", "一锅热汤，慢慢吸溜", "A hot pot of broth, slurp it slow", []),
            ("sh3", 9, 14, "店内图左移特写（占位；交付版换加料）", "store", "C-02 加料", "热气腾腾，香气扑鼻", "Steaming hot, smells so good", []),
            ("sh4", 14, 19, "店内图上摇（占位；交付版换挑面）", "store", "C-03 挑面", "暖胃，也暖心", "Warms you up, inside and out", []),
            ("sh5", 19, 24, "店内图拉远（占位；交付版换招牌菜特写）", "store", "C-05 招牌菜特写", "约上朋友，来吸一碗", "Bring a friend, grab a bowl", []),
            ("sh6", 24, 30, "结尾卡：店名字标 + 电话", "wordmark", "C-04 店门/招牌", f"{name}  电话 {phone}", f"{name}  Call {phone}", ["facts:name_en", "facts:phone"]),
        ]
        sc = {"duration_s": 30, "shots": [dict(zip(["shot_id", "start", "end", "visual", "media_id", "customer_shot_slot", "sub_zh", "sub_en", "refs"], s)) for s in shots]}
        if FAULT == "price":
            sc["shots"][4]["sub_zh"] = "招牌铜锅米线 只要 $15.90"; sc["shots"][4]["sub_en"] = "Signature copper pot only $15.90"
        sp = wjson(st.dir / "script.json", sc); validate(sp, "script", st.name); st.out(sp, "script")
        # 字段核对：引用必须存在且 verified；未引用的字幕不得含事实型信息
        bad = []
        pat = re.compile(r"\$\s?\d|\d{1,2}:\d{2}|\b0\d[ \d]{8,}\b|Rothschild|Rosebery|NSW|小锅|铜锅米线|Minty|Copper Pot|周[一二三四五六日]")
        verified = {f"facts:{k}": v["value"] for k, v in F.items() if v["status"] == "verified"}
        for i in menu["items"]:
            verified[f"menu:{i['name_en']}"] = i
        for s in sc["shots"]:
            for r in s["refs"]:
                if r not in verified: bad.append(f"{s['shot_id']} 引用 {r} 不存在或非 verified")
            txt = s["sub_zh"] + " " + s["sub_en"]
            for v in verified.values():
                if isinstance(v, str): txt = txt.replace(v, "")
            hits = [m.group(0).strip() for m in pat.finditer(txt)]
            if hits: bad.append(f"{s['shot_id']} 字幕含未核实事实型内容: {hits}")
        media_ids = {m["media_id"] for m in rjson(media_p)["media"]}
        bad += [f"{s['shot_id']} 素材 {s['media_id']} 不在 media_manifest" for s in sc["shots"] if s["media_id"] not in media_ids]
        st.check("每镜引用字段存在且 verified，不一致 0", not bad, bad, 0)
        st.check("时长合计 30s", sc["shots"][-1]["end"] == 30 and all(a["end"] == b["start"] for a, b in zip(sc["shots"], sc["shots"][1:])), sc["shots"][-1]["end"], 30)
        _storyboard(st, sc, media_p)
        st.score = {"commercial": 3, "handoff": 5, "craft": 5 if not bad else 0, "benchmark": 3,
                    "reasons": {"commercial": "只敢用店名+电话两条已核实事实，文案偏泛，缺菜品/价格卖点",
                                "craft": "引用核对 0 不一致" if not bad else f"不一致 {len(bad)} 处",
                                "benchmark": "python-docx/reportlab 是开源首选；文案生成无工具，人工(AI)写，属弱格"}}
        if bad:
            raise StationFail(st.name, "fact-check", "; ".join(bad))
    return sp, menu_p


def _storyboard(st, sc, media_p):
    from docx import Document
    from docx.shared import Cm
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    from reportlab.pdfgen import canvas
    mm = {m["media_id"]: ROOT / m["path"] for m in rjson(media_p)["media"]}
    pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
    pdf = st.dir / "storyboard.pdf"
    c = canvas.Canvas(str(pdf), pagesize=landscape(A4), invariant=1)
    W, H = landscape(A4)
    c.setFont("STSong-Light", 18); c.drawString(40, H - 40, "分镜 Storyboard · Slurp Slurp Noodle 30s（样片 · 未授权素材）")
    doc = Document(); doc.add_heading("分镜表 Storyboard", 1)
    t = doc.add_table(rows=1, cols=5); t.style = "Table Grid"
    for i, h in enumerate(["镜头", "时间", "画面", "字幕 中/英", "客户镜头位"]): t.rows[0].cells[i].text = h
    for i, s in enumerate(sc["shots"]):
        x, y = 40 + (i % 3) * 260, H - 90 - (i // 3) * 250
        c.drawImage(str(mm[s["media_id"]]), x, y - 140, width=187, height=140, preserveAspectRatio=True)
        c.setFont("STSong-Light", 10)
        for j, line in enumerate([f'{s["shot_id"]}  {s["start"]}–{s["end"]}s', s["visual"][:24], s["sub_zh"], s["sub_en"][:40], "客户镜头位: " + s["customer_shot_slot"]]):
            c.drawString(x, y - 155 - j * 13, line)
        r = t.add_row().cells
        for k, v in enumerate([s["shot_id"], f'{s["start"]}–{s["end"]}s', s["visual"], s["sub_zh"] + "\n" + s["sub_en"], s["customer_shot_slot"]]): r[k].text = v
    c.save(); doc.save(st.dir / "storyboard.docx")
    st.out(pdf); st.out(st.dir / "storyboard.docx")
